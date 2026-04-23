#!/usr/bin/env python3
"""
Merge online Excel responses + paper CSV responses into two unified CSV files.
One file per respondent group: educators and parents.

Codebook from generate_forms.gs:
  - EDU form N (1-18) → EDU_COMBOS[N-1] = [T, V]
  - PAR form N (1-27) → PAR_COMBOS[N-1] = [T, V]
"""

import csv
import re
import openpyxl
from pathlib import Path

# ── Codebook (from generate_forms.gs) ────────────────────────────────────────
EDU_COMBOS = [
    [8,5],[8,3],[9,6],[6,6],[7,4],[4,6],
    [9,3],[2,3],[7,2],[4,4],[1,2],[1,4],
    [6,5],[2,1],[5,5],[5,1],[3,2],[3,1]
]
PAR_COMBOS = [
    [2,2],[9,4],[7,5],[6,2],[4,6],[1,5],
    [8,6],[9,3],[7,4],[6,6],[2,5],[2,3],
    [4,2],[8,3],[4,5],[6,5],[3,2],[8,2],
    [5,3],[9,6],[3,3],[5,1],[3,1],[1,4],
    [7,1],[1,1],[5,6]
]

EDU_CODEBOOK = {i+1: (f"T{t}", f"V{v}") for i, (t, v) in enumerate(EDU_COMBOS)}
PAR_CODEBOOK = {i+1: (f"T{t}", f"V{v}") for i, (t, v) in enumerate(PAR_COMBOS)}

# ── File paths ────────────────────────────────────────────────────────────────
HERE = Path(__file__).parent
ONLINE_EDU  = Path("/home/dimka/Downloads/Вихователі.xlsx")
ONLINE_PAR  = Path("/home/dimka/Downloads/Батьки.xlsx")
PAPER_EDU   = HERE / "responses_educators(2).csv"
PAPER_PAR   = HERE / "responses_parents(1).csv"
OUT_EDU     = HERE / "responses_educators_all.csv"
OUT_PAR     = HERE / "responses_parents_all.csv"

# ── Unified column schemas ────────────────────────────────────────────────────
EDU_COLS = [
    "respondent_id", "questionnaire_id", "group_code", "group",
    "institution_code", "questionnaire_number", "mode", "source_label",
    "date_filled",
    "sample1_code", "sample1_format", "sample2_code", "sample2_format",
    # Block A
    "years_experience", "group_name", "education",
    "uses_video_audio_content", "use_frequency_if_yes",
    "usage_purpose_multi", "content_source_multi",
    # Block B
    "importance_precheck_1_5", "quality_signs_multi",
    "knows_ai_content", "used_ai_in_work",
    "ai_content_pedagogically_appropriate", "youtube_trust_1_5",
    # Sample 1 (text)
    "sample1_plot", "sample1_value", "sample1_language",
    "sample1_psychological_safety", "sample1_age_fit", "sample1_overall",
    "sample1_guessed_origin", "sample1_would_use",
    # Sample 2 (video)
    "sample2_plot", "sample2_value", "sample2_language",
    "sample2_psychological_safety", "sample2_age_fit", "sample2_overall",
    "sample2_guessed_origin", "sample2_would_use",
    # Open
    "open_q13_main_risk_ai", "open_q14_quality_signs",
    # Extras from paper
    "usage_purpose_other_text", "content_source_other_text",
    "quality_signs_other_text", "notes",
]

PAR_COLS = [
    "respondent_id", "questionnaire_id", "group_code", "group",
    "institution_code", "questionnaire_number", "mode", "source_label",
    "date_filled",
    "sample1_code", "sample1_format", "sample2_code", "sample2_format",
    # Block A
    "child_age_years", "child_gender", "preschool_children_in_family",
    "weekday_video_time", "primary_device", "content_chooser",
    "co_viewing_frequency",
    # Block B
    "precheck_frequency", "selection_criteria_multi",
    "content_quality_concern_1_5", "knows_ai_content",
    "has_given_ai_content", "negative_reaction_seen", "youtube_trust_1_5",
    # Sample 1 (text)
    "sample1_likes", "sample1_positive_value", "sample1_safe",
    "sample1_language_quality", "sample1_age_fit", "sample1_overall",
    "sample1_guessed_origin", "sample1_would_show",
    # Sample 2 (video)
    "sample2_likes", "sample2_positive_value", "sample2_safe",
    "sample2_language_quality", "sample2_age_fit", "sample2_overall",
    "sample2_guessed_origin", "sample2_would_show",
    # Open
    "open_q14_most_important", "open_q15_what_avoid",
    # Extras from paper
    "selection_criteria_other_text", "notes",
]


def parse_sheet_num(sheet_name):
    """'Відповіді форми (7)' → 7"""
    m = re.search(r'\((\d+)\)', sheet_name)
    return int(m.group(1)) if m else None


ARTIFACT_RE = re.compile(r'Анкета\s+([ВБ]-\d+-T(\d+)V(\d+))', re.IGNORECASE)


def detect_sheet_codes(ws, fallback_questionnaire_num, codebook, group_prefix):
    """
    Look for an artifact row of the form 'Анкета В-04-T6V6'.
    Returns (questionnaire_id, t_code, v_code).
    Falls back to codebook[fallback_questionnaire_num] if no artifact found.
    """
    for row in ws.iter_rows(min_row=2, values_only=True):
        non_null = [v for v in row if v is not None]
        if len(non_null) == 1 and isinstance(non_null[0], str):
            m = ARTIFACT_RE.search(non_null[0])
            if m:
                questionnaire_id = m.group(1)
                t_code = f"T{m.group(2)}"
                v_code = f"V{m.group(3)}"
                return questionnaire_id, t_code, v_code
    # No artifact — use sheet number fallback
    t_code, v_code = codebook[fallback_questionnaire_num]
    questionnaire_id = f"{group_prefix}-{fallback_questionnaire_num:02d}-{t_code}{v_code}"
    return questionnaire_id, t_code, v_code


def read_online_educators(path):
    wb = openpyxl.load_workbook(path)
    rows = []
    counter = 1
    for sheet_name in sorted(wb.sheetnames, key=lambda s: parse_sheet_num(s) or 0):
        form_num = parse_sheet_num(sheet_name)
        if form_num is None:
            continue
        ws = wb[sheet_name]
        questionnaire_id, t_code, v_code = detect_sheet_codes(ws, form_num, EDU_CODEBOOK, "В")
        # Extract the questionnaire number from the questionnaire_id (В-04-T6V6 → 4)
        qnum_match = re.search(r'[ВБ]-(\d+)-', questionnaire_id)
        qnum = int(qnum_match.group(1)) if qnum_match else form_num
        first = True
        for row in ws.iter_rows(values_only=True):
            if first:
                first = False
                continue  # skip header
            if all(v is None for v in row):
                continue
            # Skip artifact rows (< 5 non-null values)
            if sum(1 for v in row if v is not None) < 5:
                continue
            r = dict.fromkeys(EDU_COLS, "")
            r["respondent_id"]      = f"E-online-{counter:03d}"
            r["questionnaire_id"]   = questionnaire_id
            r["group_code"]         = "В"
            r["group"]              = "educator"
            r["institution_code"]   = "online"
            r["questionnaire_number"] = str(qnum)
            r["mode"]               = "google_form"
            r["source_label"]       = "Google Forms"
            r["date_filled"]        = str(row[0]) if row[0] else ""
            r["sample1_code"]       = t_code
            r["sample1_format"]     = "text"
            r["sample2_code"]       = v_code
            r["sample2_format"]     = "video"
            # Block A
            r["years_experience"]              = _v(row, 1)
            r["group_name"]                    = _v(row, 2)
            r["education"]                     = _v(row, 3)
            r["uses_video_audio_content"]      = _v(row, 4)
            r["use_frequency_if_yes"]          = _v(row, 5)
            r["usage_purpose_multi"]           = _v(row, 6)
            r["content_source_multi"]          = _v(row, 7)
            # Block B
            r["importance_precheck_1_5"]       = _v(row, 8)
            r["quality_signs_multi"]           = _v(row, 9)
            r["knows_ai_content"]              = _v(row, 10)
            r["used_ai_in_work"]               = _v(row, 11)
            r["ai_content_pedagogically_appropriate"] = _v(row, 12)
            r["youtube_trust_1_5"]             = _v(row, 13)
            # Sample 1 (text)
            r["sample1_plot"]                  = _v(row, 14)
            r["sample1_value"]                 = _v(row, 15)
            r["sample1_language"]              = _v(row, 16)
            r["sample1_psychological_safety"]  = _v(row, 17)
            r["sample1_age_fit"]               = _v(row, 18)
            r["sample1_overall"]               = _v(row, 19)
            r["sample1_guessed_origin"]        = _v(row, 20)
            r["sample1_would_use"]             = _v(row, 21)
            # Sample 2 (video)
            r["sample2_plot"]                  = _v(row, 22)
            r["sample2_value"]                 = _v(row, 23)
            r["sample2_language"]              = _v(row, 24)
            r["sample2_psychological_safety"]  = _v(row, 25)
            r["sample2_age_fit"]               = _v(row, 26)
            r["sample2_overall"]               = _v(row, 27)
            r["sample2_guessed_origin"]        = _v(row, 28)
            r["sample2_would_use"]             = _v(row, 29)
            # Open
            r["open_q13_main_risk_ai"]         = _v(row, 30)
            r["open_q14_quality_signs"]        = _v(row, 31)
            rows.append(r)
            counter += 1
    print(f"Online educators: {len(rows)} rows from {path.name}")
    return rows


def read_online_parents(path):
    wb = openpyxl.load_workbook(path)
    rows = []
    counter = 1
    for sheet_name in sorted(wb.sheetnames, key=lambda s: parse_sheet_num(s) or 0):
        form_num = parse_sheet_num(sheet_name)
        if form_num is None:
            continue
        ws = wb[sheet_name]
        questionnaire_id, t_code, v_code = detect_sheet_codes(ws, form_num, PAR_CODEBOOK, "Б")
        qnum_match = re.search(r'[ВБ]-(\d+)-', questionnaire_id)
        qnum = int(qnum_match.group(1)) if qnum_match else form_num
        first = True
        for row in ws.iter_rows(values_only=True):
            if first:
                first = False
                continue
            if all(v is None for v in row):
                continue
            # Skip artifact rows
            if sum(1 for v in row if v is not None) < 5:
                continue
            r = dict.fromkeys(PAR_COLS, "")
            r["respondent_id"]      = f"P-online-{counter:03d}"
            r["questionnaire_id"]   = questionnaire_id
            r["group_code"]         = "Б"
            r["group"]              = "parent"
            r["institution_code"]   = "online"
            r["questionnaire_number"] = str(qnum)
            r["mode"]               = "google_form"
            r["source_label"]       = "Google Forms"
            r["date_filled"]        = str(row[0]) if row[0] else ""
            r["sample1_code"]       = t_code
            r["sample1_format"]     = "text"
            r["sample2_code"]       = v_code
            r["sample2_format"]     = "video"
            # Block A
            r["child_age_years"]               = _v(row, 1)
            r["child_gender"]                  = _v(row, 2)
            r["preschool_children_in_family"]  = _v(row, 3)
            r["weekday_video_time"]            = _v(row, 4)
            r["primary_device"]                = _v(row, 5)
            r["content_chooser"]               = _v(row, 6)
            r["co_viewing_frequency"]          = _v(row, 7)
            # Block B
            r["precheck_frequency"]            = _v(row, 8)
            r["selection_criteria_multi"]      = _v(row, 9)
            r["content_quality_concern_1_5"]   = _v(row, 10)
            r["knows_ai_content"]              = _v(row, 11)
            r["has_given_ai_content"]          = _v(row, 12)
            r["negative_reaction_seen"]        = _v(row, 13)
            r["youtube_trust_1_5"]             = _v(row, 14)
            # Sample 1 (text)
            r["sample1_likes"]                 = _v(row, 15)
            r["sample1_positive_value"]        = _v(row, 16)
            r["sample1_safe"]                  = _v(row, 17)
            r["sample1_language_quality"]      = _v(row, 18)
            r["sample1_age_fit"]               = _v(row, 19)
            r["sample1_overall"]               = _v(row, 20)
            r["sample1_guessed_origin"]        = _v(row, 21)
            r["sample1_would_show"]            = _v(row, 22)
            # Sample 2 (video)
            r["sample2_likes"]                 = _v(row, 23)
            r["sample2_positive_value"]        = _v(row, 24)
            r["sample2_safe"]                  = _v(row, 25)
            r["sample2_language_quality"]      = _v(row, 26)
            r["sample2_age_fit"]               = _v(row, 27)
            r["sample2_overall"]               = _v(row, 28)
            r["sample2_guessed_origin"]        = _v(row, 29)
            r["sample2_would_show"]            = _v(row, 30)
            # Open
            r["open_q14_most_important"]       = _v(row, 31)
            r["open_q15_what_avoid"]           = _v(row, 32)
            rows.append(r)
            counter += 1
    print(f"Online parents: {len(rows)} rows from {path.name}")
    return rows


def read_paper_educators(path):
    rows = []
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for raw in reader:
            r = dict.fromkeys(EDU_COLS, "")
            for col in EDU_COLS:
                if col in raw:
                    r[col] = raw[col]
            # date_filled not in paper CSV — leave blank
            r.setdefault("date_filled", "")
            rows.append(r)
    print(f"Paper educators: {len(rows)} rows from {path.name}")
    return rows


def read_paper_parents(path):
    rows = []
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for raw in reader:
            r = dict.fromkeys(PAR_COLS, "")
            for col in PAR_COLS:
                if col in raw:
                    r[col] = raw[col]
            r.setdefault("date_filled", "")
            rows.append(r)
    print(f"Paper parents: {len(rows)} rows from {path.name}")
    return rows


def write_csv(path, cols, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    print(f"Written {len(rows)} rows → {path.name}")


def write_xlsx(path, cols, rows):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Дані"

    # Header style
    header_fill = PatternFill("solid", fgColor="1A56A0")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_border = Border(
        bottom=Side(style="thin", color="FFFFFF"),
        right=Side(style="thin", color="FFFFFF"),
    )
    # Alternate row fill
    row_fill_alt = PatternFill("solid", fgColor="EFF4FC")

    # Write header
    for col_idx, col_name in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border

    ws.row_dimensions[1].height = 42

    # Write data
    for row_idx, row in enumerate(rows, start=2):
        fill = row_fill_alt if row_idx % 2 == 0 else None
        for col_idx, col_name in enumerate(cols, start=1):
            val = row.get(col_name, "")
            # Try to keep numbers as numbers
            if val != "":
                try:
                    val = float(val)
                    if val == int(val):
                        val = int(val)
                except (ValueError, TypeError):
                    pass
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fill:
                cell.fill = fill

    # Auto-fit column widths (capped)
    for col_idx, col_name in enumerate(cols, start=1):
        letter = get_column_letter(col_idx)
        # Sample first 30 rows to estimate width
        max_len = len(col_name)
        for row_idx in range(2, min(len(rows) + 2, 32)):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val:
                max_len = max(max_len, min(len(str(val)), 60))
        ws.column_dimensions[letter].width = max(8, min(max_len + 2, 38))

    # Freeze header row
    ws.freeze_panes = "A2"

    # Auto-filter
    ws.auto_filter.ref = ws.dimensions

    wb.save(path)
    print(f"Written {len(rows)} rows → {path.name}")


def _v(row, idx):
    """Safe value getter from tuple row."""
    try:
        val = row[idx]
        return "" if val is None else str(val)
    except IndexError:
        return ""


if __name__ == "__main__":
    online_edu  = read_online_educators(ONLINE_EDU)
    paper_edu   = read_paper_educators(PAPER_EDU)
    all_edu     = paper_edu + online_edu
    write_csv(OUT_EDU, EDU_COLS, all_edu)
    write_xlsx(OUT_EDU.with_suffix(".xlsx"), EDU_COLS, all_edu)

    online_par  = read_online_parents(ONLINE_PAR)
    paper_par   = read_paper_parents(PAPER_PAR)
    all_par     = paper_par + online_par
    write_csv(OUT_PAR, PAR_COLS, all_par)
    write_xlsx(OUT_PAR.with_suffix(".xlsx"), PAR_COLS, all_par)

    print(f"\nDone. Educators: {len(all_edu)} total | Parents: {len(all_par)} total")
